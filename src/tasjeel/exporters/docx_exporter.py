from .base_exporter import BaseExporter
from typing import Any, Dict
import time

RTL_LANGUAGES = ["Arabic", "Persian", "Urdu"]
DOCUMENT_TITLE = "Voice Transcription"
HEADING_LEVEL = 0

class DocxExporter(BaseExporter):
    KEY = "docx"
    LABEL = "Microsoft Word (.docx)"
    FILE_FILTER = "Word documents (*.docx)"

    @classmethod
    def check_dependencies(cls) -> bool:
        try:
            import docx
            return True
        except ImportError:
            return False

    def export(self, file_path: str, data: Dict[str, Any]) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        doc.add_heading(DOCUMENT_TITLE, HEADING_LEVEL)
        doc.add_paragraph(f"Generated on: {time.strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Language: {data.get('language', 'Unknown')}")
        doc.add_paragraph("")
        
        paragraph = doc.add_paragraph(data.get("text", ""))
        if data.get("language") in RTL_LANGUAGES:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
        doc.save(file_path)
